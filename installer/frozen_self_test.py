"""Optional PyInstaller runtime smoke test, before application initialization.

TikTokBatchMVP.exe --self-test <report.json>
Checks bundled code/assets only. Does not read preferences, cookies or API keys,
open a browser/window, or send any network requests.
"""
import sys


if "--self-test" in sys.argv:
    import importlib
    import json
    import platform
    from pathlib import Path

    argument_index = sys.argv.index("--self-test")
    if argument_index + 1 >= len(sys.argv):
        raise SystemExit(2)
    report_path = Path(sys.argv[argument_index + 1]).resolve()
    # The build knows which version it is producing; the bundle has to agree.
    # Passing it in turns "VERSION got left out of the package" into a build
    # failure instead of an app that silently reports 0.0.0 forever.
    expected_version = None
    if "--expect-version" in sys.argv:
        version_index = sys.argv.index("--expect-version")
        if version_index + 1 >= len(sys.argv):
            raise SystemExit(2)
        expected_version = sys.argv[version_index + 1].strip()
    asset_root = (
        Path(sys._MEIPASS)
        if getattr(sys, "frozen", False)
        else Path(__file__).resolve().parent.parent / "outputs" / "TikTokBatchMVP"
    )
    checks = []

    def check(name, operation):
        try:
            detail = operation()
            checks.append({"name": name, "ok": True, "detail": str(detail or "OK")})
        except Exception as exc:
            checks.append({"name": name, "ok": False, "error": f"{type(exc).__name__}: {exc}"})

    def require_file(path):
        if not path.is_file() or path.stat().st_size == 0:
            raise FileNotFoundError(f"Bundled resource missing or empty: {path.name}")
        return path.name

    for module_name in (
        "yt_dlp",
        "yt_dlp.networking.impersonate",
        "yt_dlp.networking._curlcffi",
        "curl_cffi.requests",
        "playwright.sync_api",
        "webview.platforms.edgechromium",
        "docx",
        "requests",
        "session_store",
        "profile_pagination",
    ):
        check(f"import:{module_name}", lambda name=module_name: importlib.import_module(name).__name__)

    for relative_path in ("ui/index.html", "ui/blank.html", "ui/tiktok-logo.png", "VERSION"):
        check(f"asset:{relative_path}", lambda relative=relative_path: require_file(asset_root / relative))

    def check_bundled_version():
        bundled = (asset_root / "VERSION").read_text(encoding="utf-8").strip()
        if not bundled or not bundled[0].isdigit():
            raise ValueError(f"Bundled VERSION is not a version number: {bundled!r}")
        if expected_version and bundled != expected_version:
            raise ValueError(f"Bundled VERSION is {bundled!r} but this build produces {expected_version!r}")
        return f"{bundled} (matches build)" if expected_version else bundled

    check("version:matches-build", check_bundled_version)

    def check_playwright_driver():
        from playwright._impl._driver import compute_driver_executable
        from playwright.sync_api import sync_playwright

        node, cli = compute_driver_executable()
        require_file(Path(node))
        require_file(Path(cli))
        # Starts/stops the bundled Node driver only; no browser gets launched.
        with sync_playwright() as playwright:
            if playwright.chromium.name != "chromium":
                raise RuntimeError("Playwright driver did not initialize Chromium support")
        return "Node driver initialized without launching a browser"

    check("playwright:driver", check_playwright_driver)

    def check_docx_template():
        from docx import Document
        from io import BytesIO

        document = Document()
        document.add_paragraph("Bundled document template check")
        content = BytesIO()
        document.save(content)
        content.seek(0)
        if Document(content).paragraphs[0].text != "Bundled document template check":
            raise RuntimeError("Word document round trip failed")
        return "Default template and DOCX save/load passed in memory"

    check("docx:template", check_docx_template)

    def check_webview_loader():
        import webview

        library_root = Path(webview.__file__).parent / "lib"
        require_file(library_root / "Microsoft.Web.WebView2.Core.dll")
        machine = platform.machine().lower()
        architecture = "arm64" if machine in ("arm64", "aarch64") else ("x64" if sys.maxsize > 2**32 else "x86")
        return require_file(library_root / "runtimes" / f"win-{architecture}" / "native" / "WebView2Loader.dll")

    check("webview:loader", check_webview_loader)
    report = {"ok": all(item["ok"] for item in checks), "frozen": bool(getattr(sys, "frozen", False)), "checks": checks}
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    raise SystemExit(0 if report["ok"] else 1)
